# REFERENCE implementation for the end-of-run Excel report (see CLAUDE.md "Excel report").
# Copy into the run's cache/, swap ROWS/title/subtitle/OUT for the current run's data.
# The ROWS below are FICTIONAL placeholders demonstrating one row of each case —
# never let placeholder values leak into a real report.
# Native cell hyperlinks (no formulas) so no recalc pass is required;
# PDF links carry #page=N so browsers open the report at the evidence page.
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "Scope-1-2-Emissions-Report.xlsx"

# company, website, outcome, year, s1, s2m, s2l, s12comb, confidence, note, link_label, link_url, needs_review
ROWS = [
    # Case 1: PDF with exact table figures, both Scope 2 bases -> High, link opens at evidence page
    ("Acme Corp", "https://www.acme.example/", "pdf", 2025, 12345, 2345, 3456, None, "High",
     "Sustainability Report 2025 table; both Scope 2 bases reported.",
     "Sustainability Report 2025 · p.42",
     "https://www.acme.example/reports/sustainability-2025.pdf#page=42", False),
    # Case 2: figures visible on a webpage but interpretation applied -> Medium, link is the live page
    ("Globex Industries", "https://www.globex.example/", "screenshot", 2024, 500000, None, 700000, None, "Medium",
     "On-page values coarsely rounded; Scope 2 basis not stated — inferred location.",
     "Live page — Environmental data",
     "https://www.globex.example/sustainability/environmental-data", False),
    # Case 3: combined-only reporter -> value in the combined column ONLY, split columns empty
    ("Initech LLC", "https://www.initech.example/", "pdf", 2024, None, None, None, 9876.5, "High",
     "Reported combined only (no Scope 1/2 split; basis within total not stated).",
     "ESG Report 2024 · p.7",
     "https://www.initech.example/esg-report-2024.pdf#page=7", False),
    # Case 4: needs_review — e.g. rebrand/parent-sourced values. The review flag (last
    # field, from queue status) forces Confidence to "Needs review", overriding any rating.
    ("Umbrella Holdings", "https://www.umbrella.example/", "pdf", 2024, 55555, None, 66666, None, "High",
     "Rebranded to Parasol (identity needs review); Scope 2 basis inferred.",
     "Parasol Sustainability Report 2024 · p.12",
     "https://www.parasol.example/reports/sustainability-2024.pdf#page=12", True),
    # Case 5: nothing published -> muted grey row, dash confidence, link is the company site
    ("Hooli Inc", "https://www.hooli.example/", "none", None, None, None, None, None, "—",
     "No ESG/GHG disclosure found (private company; site verified live).",
     "Company site (no disclosure)", "https://www.hooli.example/", False),
]

NAVY = "1F3864"
BAND_FILL = PatternFill("solid", fgColor="F2F2F2")      # zebra banding (the only row fill)
CONF_COLOR = {"High": "1E7B34", "Medium": "B45F06", "Needs review": "C00000", "—": "808080"}
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "Scope 1-2 Results"
ws.sheet_view.showGridLines = False

HEADERS = ["Company", "Website", "Outcome", "Report year", "Scope 1 (tCO2e)",
           "Scope 2 market-based (tCO2e)", "Scope 2 location-based (tCO2e)",
           "Scope 1+2 combined (tCO2e)", "Confidence", "Note",
           "Source (click to open at page)"]
WIDTHS = [26, 30, 11, 11, 15, 15, 15, 15, 14, 62, 38]
CENTER_COLS = (3, 4, 9)     # outcome, year, confidence
NUM_COLS = (5, 6, 7, 8)     # emission values
NOTE_COL, SRC_COL = 10, 11

ws["A1"] = "Scope 1 & 2 Emissions — <Batch Name>"
ws["A1"].font = Font(name="Arial", bold=True, size=14, color=NAVY)
ws["A2"] = "Run <YYYY-MM-DD>_run-<nn> · <N> companies · all emission values in tCO2e (latest reported year) · 'Needs review' rows await a human ruling (reasons in queue.csv)"
ws["A2"].font = Font(name="Arial", size=9, color="808080")

HDR = 4
for c, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=1):
    cell = ws.cell(row=HDR, column=c, value=h)
    cell.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[HDR].height = 30

for i, (name, site, outcome, year, s1, s2m, s2l, s12, conf, note, label, url, review) in enumerate(ROWS):
    r = HDR + 1 + i
    fill = BAND_FILL if i % 2 else None
    # muted = no emission values at all (covers outcome=none AND collected-report-but-
    # nothing-published rows) — keyed off the values, not the outcome
    muted = all(v is None for v in (s1, s2m, s2l, s12))
    # one signal per row: a pending human ruling overrides any data-quality rating
    if review:
        conf = "Needs review"
    base = dict(name="Arial", size=10, color="A6A6A6" if muted else "000000")

    vals = [name, site, outcome, year, s1, s2m, s2l, s12, conf, note, label]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=c, value="—" if (v is None and c in (3, 4)) else v)
        cell.font = Font(**base)
        cell.border = BORDER
        if fill:
            cell.fill = fill
        if c in CENTER_COLS:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c in NUM_COLS:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right", vertical="center")
        else:
            cell.alignment = Alignment(vertical="center", wrap_text=(c == NOTE_COL))

    ws.cell(row=r, column=1).font = Font(name="Arial", size=10, bold=True,
                                         color="A6A6A6" if muted else "000000")
    if conf in CONF_COLOR:
        ws.cell(row=r, column=9).font = Font(name="Arial", size=10, bold=conf != "—",
                                             color=CONF_COLOR[conf])
    link_font = Font(name="Arial", size=10, underline="single",
                     color="A6A6A6" if muted else "0563C1")
    site_cell = ws.cell(row=r, column=2)
    site_cell.hyperlink = site
    site_cell.font = link_font
    src_cell = ws.cell(row=r, column=SRC_COL)
    src_cell.hyperlink = url
    src_cell.font = link_font

last = HDR + len(ROWS)
ws.freeze_panes = f"A{HDR + 1}"
ws.auto_filter.ref = f"A{HDR}:K{last}"

legend = [
    ("Confidence — one signal per row (data age is NOT a factor; see Report year):", True),
    ("  High — exact figures lifted straight from the company's own report/page; no interpretation needed.", False),
    ("  Medium — interpretation applied: Scope 2 basis inferred, unit conversion/coarse rounding, or values read off a chart rather than a table.", False),
    ("  Needs review — a human ruling is pending (values sourced from a parent/child company's report, identity unconfirmed, blocked file, or doubtful no-disclosure); overrides any rating. Reason in queue.csv.", False),
    ("  — (dash) — no figures published and no ruling pending.", False),
    ("Grey text rows = no absolute Scope 1/2 figures published (with or without a collected report).", False),
    ("Scope 1+2 combined is filled only when a company publishes a merged total without the split.", False),
    ("Source column: PDF links open the company's hosted report in your browser at the cited page (#page=N); 'Live page' links open the webpage that was screenshotted.", False),
    ("Original evidence files (PDFs/screenshots) are stored under companies\\<Name>\\ in this run folder; raw tracker: results.csv.", False),
]
for j, (txt, bold) in enumerate(legend):
    cell = ws.cell(row=last + 2 + j, column=1, value=txt)
    cell.font = Font(name="Arial", size=9, bold=bold, color="595959")

wb.save(OUT)
print("saved", OUT)
